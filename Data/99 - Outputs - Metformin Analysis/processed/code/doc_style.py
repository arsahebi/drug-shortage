# %%
"""
Shared Word styling for the Metformin analysis documents.

Black-and-white academic style matching
Data/99 - Outputs - Text Analysis/eval/483_Worked_Example_FEI3003342394_obs3.docx:
Times New Roman, no accent colours, thin ruled tables with a grey header row,
shaded single-cell callout boxes.

Used by build_changelog_doc.py and build_figures_doc.py so both documents stay
visually identical.
"""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

FONT     = "Times New Roman"
BLACK    = RGBColor(0x00, 0x00, 0x00)
HDR_FILL = "D9D9D9"
BOX_FILL = "F2F2F2"


def new_document() -> Document:
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(11)
    normal.font.color.rgb = BLACK
    normal.paragraph_format.space_after = Pt(8)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)

    for name, size in [("Title", 20), ("Heading 1", 15),
                       ("Heading 2", 13), ("Heading 3", 11)]:
        st = doc.styles[name]
        st.font.name = FONT
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.italic = False
        st.font.color.rgb = BLACK
        if st.element.rPr is not None and st.element.rPr.rFonts is not None:
            st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)

    for s in doc.sections:
        s.left_margin = s.right_margin = Inches(1)
        s.top_margin = s.bottom_margin = Inches(1)
    return doc


def shade(cell, fill: str) -> None:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shd)


def set_borders(tbl, sz: int = 4) -> None:
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(sz))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000")
        borders.append(el)
    tbl._tbl.tblPr.append(borders)


def _style_run(run, size, bold=False, italic=False):
    run.bold, run.italic = bold, italic
    run.font.size = Pt(size)
    run.font.name = FONT
    run.font.color.rgb = BLACK
    return run


def p(doc, text="", bold=False, italic=False, size=11, align=None):
    par = doc.add_paragraph()
    if align is not None:
        par.alignment = align
    if text:
        _style_run(par.add_run(text), size, bold, italic)
    return par


def bullet(doc, text, size=11):
    par = doc.add_paragraph(style="List Bullet")
    _style_run(par.add_run(text), size)
    return par


def table(doc, headers, rows, size=9.5, widths=None, align_right_from=None):
    """Ruled table. Wrap a cell value in ** ** to bold it."""
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_borders(t)
    for i, htxt in enumerate(headers):
        c = t.rows[0].cells[i]
        shade(c, HDR_FILL)
        c.text = ""
        par = c.paragraphs[0]
        par.paragraph_format.space_after = Pt(2)
        _style_run(par.add_run(htxt), size, bold=True)
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            par = cells[i].paragraphs[0]
            par.paragraph_format.space_after = Pt(2)
            if align_right_from is not None and i >= align_right_from:
                par.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            txt = str(val)
            bold = txt.startswith("**") and txt.endswith("**") and len(txt) > 4
            if bold:
                txt = txt[2:-2]
            _style_run(par.add_run(txt), size, bold=bold)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Inches(w)
    doc.add_paragraph()
    return t


def box(doc, lines, size=10):
    """Single-cell shaded callout. lines = [(bold_label, text), ...]."""
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_borders(t, sz=6)
    c = t.rows[0].cells[0]
    shade(c, BOX_FILL)
    c.text = ""
    par = c.paragraphs[0]
    for i, (label, txt) in enumerate(lines):
        if i:
            par = c.add_paragraph()
        par.paragraph_format.space_after = Pt(3)
        if label:
            _style_run(par.add_run(f"{label}  "), size, bold=True)
        _style_run(par.add_run(txt), size)
    doc.add_paragraph()
    return t


def figure(doc, png_path: Path, caption: str, width=6.0):
    par = doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par.paragraph_format.space_after = Pt(4)
    if Path(png_path).exists():
        par.add_run().add_picture(str(png_path), width=Inches(width))
    else:
        _style_run(par.add_run("(figure not found)"), 10, italic=True)
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _style_run(cap.add_run(caption), 9, italic=True)


def rule(doc):
    par = doc.add_paragraph()
    par.paragraph_format.space_after = Pt(6)
    pPr = par._p.get_or_add_pPr()
    b = OxmlElement("w:pBdr")
    bot = OxmlElement("w:bottom")
    bot.set(qn("w:val"), "single")
    bot.set(qn("w:sz"), "6")
    bot.set(qn("w:space"), "1")
    bot.set(qn("w:color"), "000000")
    b.append(bot)
    pPr.append(b)
    return par
# %%
