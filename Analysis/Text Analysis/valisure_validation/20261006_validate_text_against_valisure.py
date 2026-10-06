"""Exploratory text/Valisure validation across both complete DoD workbooks.

Read README.md for the frozen first-pass design and limitations. Source files
are never rewritten. No LLM calls are made. CSVs preserve raw assay values and
source row provenance; the report does not interpret score penalties as harm.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import re
from pathlib import Path

import numpy as np
import pandas as pd
import statsmodels.api as sm
from openpyxl.utils import get_column_letter
from scipy.stats import t as student_t
from sklearn.compose import ColumnTransformer
from sklearn.linear_model import LogisticRegression
from sklearn.metrics import brier_score_loss, roc_auc_score
from sklearn.model_selection import GroupKFold
from sklearn.pipeline import make_pipeline
from sklearn.preprocessing import OneHotEncoder, StandardScaler
from statsmodels.stats.multitest import multipletests

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[2]
DATA = ROOT / "Data"
TEXT = ROOT / "Analysis/Text Analysis"
VAL = DATA / "08 - Valisure/raw"
FILES = {
    "testing": VAL / "Testing Data_DoD First 13 Drug Scores with ANDAs & NDCs.xlsx",
    "scoring": VAL / "Discrete Scoring_DoD First 13 Drug Scores with ANDAs & NDCs.xlsx",
    "march": VAL / "FEIs_March 2026.xlsx",
    "propublica": DATA / "19 - ProPublica/raw/ndc_fei.csv",
    "dailymed": DATA / "17 - NDC-FEI Linkage/processed/all_daily_med.csv",
    "intake": VAL / "Valisure_2024_raw_prices_20260728_f1-and-formulation_20260813.xlsx",
    "observations": TEXT / "step01_redica_483_obs_llm_signals_anthropic_claudesonnet5_v2.csv",
    "grades": TEXT / "vai_signal_validation/outputs/fei_ae_panel_inspection_centered.parquet",
}
ASSOCIATIONS = [
    ("dissolution_penalty", "screen_dissolution_share", "Dissolution screen and dissolution penalty"),
    ("chemical_penalty", "screen_chemical_share", "Chemical screen and chemical penalty"),
    ("quality_loss", "patient_risk_share", "LLM patient-risk flag and score loss"),
    ("quality_loss", "data_integrity_share", "LLM data-integrity flag and score loss"),
    ("quality_loss", "lab_share", "LLM laboratory system and score loss"),
    ("dissolution_percentile", "screen_dissolution_share", "Dissolution screen and raw difference factor"),
]
STERILE = re.compile(r"steril|aseptic|inject|parenteral|\bvials?\b|lyophili|media fill|endotoxin|cleanroom|ophthalmic", re.I)
ORAL = re.compile(r"tablet|capsule|oral solid|compress|granulat|blend|coating|encapsul", re.I)
DISSOLUTION = re.compile(r"dissolution", re.I)
CHEMICAL = re.compile(r"impurit|nitrosamin|\bNDMA\b|\bDMF\b|dimethylformamide|residual solvent|heavy metals?|elemental contamin|arsenic|thallium", re.I)
LLM_FIELDS = ["patient_risk_share", "data_integrity_share", "lab_share"]


def clean(x) -> str:
    if pd.isna(x):
        return ""
    return re.sub(r"\s+", " ", str(x)).strip()


def api_key(x) -> str:
    return re.sub(r"\s*;\s*", "; ", clean(x).lower())


def ndc(x, product=False) -> str:
    s = clean(x)
    if re.fullmatch(r"\d+\.0", s):
        s = s[:-2]
    p = s.split("-")
    if len(p) in (2, 3) and all(z.isdigit() for z in p):
        if len(p[0]) > 5 or len(p[1]) > 4:
            return ""
        base = p[0].zfill(5) + p[1].zfill(4)
        if not product and len(p) == 3 and len(p[2]) > 2:
            return ""
        return base if product else base + (p[2].zfill(2) if len(p) == 3 else "")
    if s.isdigit() and len(s) in (9, 11):
        return s[:9] if product else s
    return ""  # Bare 10-digit identifiers are ambiguous; do not guess padding.


def fei_key(x) -> str:
    s = clean(x)
    return re.sub(r"\.0$", "", s) if re.fullmatch(r"\d+(?:\.0)?", s) else ""


def assay_value(x) -> tuple[str, float, float]:
    s = clean(x)
    if not s or s.upper() in {"--", "-", "N/A", "NA", "NAN"}:
        return "not_recorded", np.nan, np.nan
    if s.upper() in {"ND", "N/D", "NOT DETECTED", "<LOQ", "LOQ", "<LOD", "LOD", "BLOQ"}:
        return "below_unspecified_limit", np.nan, np.nan
    if s.startswith("<"):
        try:
            return "below_recorded_limit", np.nan, float(s[1:].strip())
        except ValueError:
            return "unparsed", np.nan, np.nan
    try:
        return "quantified", float(s), np.nan
    except ValueError:
        return "unparsed", np.nan, np.nan


def read_samples() -> tuple[pd.DataFrame, pd.DataFrame]:
    samples, assays = [], []
    books = pd.read_excel(FILES["testing"], sheet_name=None, header=1)
    excluded = {"Sample ID", "API", "Labeler", "DoD Drug Score", "NDC", "Form Short Name", "Strength", "Application Number", "Lot Number", "Expiration Date", "Elemental", "Carcinogens", "Dissolution"}
    for sheet, frame in books.items():
        frame.columns = [clean(c) for c in frame]
        frame["API"] = frame["API"].ffill()
        for idx, row in frame[frame["Sample ID"].notna()].iterrows():
            sid = clean(row["Sample ID"])
            rec = {
                "sample_id": sid, "source_sheet": sheet, "source_row": int(idx) + 3,
                "api": api_key(row["API"]), "labeler": clean(row["Labeler"]),
                "ndc_raw": clean(row["NDC"]), "ndc11": ndc(row["NDC"]),
                "ndc9": ndc(row["NDC"], True), "form": clean(row["Form Short Name"]).upper(),
                "strength": clean(row["Strength"]).upper(),
                "application": clean(row["Application Number"]), "lot": clean(row["Lot Number"]),
                "expiration_raw": clean(row["Expiration Date"]),
                "recorded_drug_score": pd.to_numeric(row["DoD Drug Score"], errors="coerce"),
            }
            rec["product"] = " | ".join(rec[k] for k in ["api", "form", "strength"])
            rec["ndc_parse_status"] = "package_and_product_valid" if len(rec["ndc11"]) == 11 else ("product_only_package_invalid" if len(rec["ndc9"]) == 9 else "invalid_product")
            rec["route_screen"] = "oral_solid" if re.search(r"TABLET|CAPSULE", rec["form"]) else ("injectable" if re.search(r"INJECT|\bINJ\b", rec["form"]) else "unspecified")
            samples.append(rec)
            for col in frame:
                if col in excluded or col.startswith("Unnamed"):
                    continue
                status, value, limit = assay_value(row[col])
                assays.append({"sample_id": sid, "api": rec["api"], "assay": col, "raw_value": clean(row[col]), "status": status, "value": value, "recorded_upper_limit": limit, "source_sheet": sheet, "source_row": int(idx) + 3,"source_cell":f"{get_column_letter(frame.columns.get_loc(col)+1)}{int(idx)+3}"})
    s = pd.DataFrame(samples)
    assert not s.sample_id.duplicated().any(), "Repeated sample IDs need adjudication before merging"
    assert s.ndc9.str.len().eq(9).all(), "Invalid product NDC in source; inspect before proceeding"
    a = pd.DataFrame(assays)
    d = a[a.assay.eq("Difference Factor")].set_index("sample_id")
    s["difference_factor"] = s.sample_id.map(d.value)
    chem = a[~a.assay.eq("Difference Factor") & ~a.status.eq("not_recorded")].sample_id.unique()
    s["has_chemical_measurement"] = s.sample_id.isin(chem)
    s["has_dissolution_measurement"] = s.difference_factor.notna()
    return s, a


def read_scores() -> tuple[pd.DataFrame, pd.DataFrame]:
    records, expanded = [], []
    for sheet, d in pd.read_excel(FILES["scoring"], sheet_name=None, header=0).items():
        d.columns = [clean(c) for c in d]
        cols = list(d)
        score_col = next(c for c in cols if c.lower() == "score")
        app_idx, ndc_idx = cols.index("Application Number"), cols.index("NDCs")
        components = [c for c in cols[:app_idx] if c not in {"API", "Company", score_col}]
        for idx, row in d[d.Company.notna()].iterrows():
            cid = f"{sheet}:{int(idx)+2}"
            vals = {}
            for c in components:
                raw = clean(row[c])
                if raw == "--":
                    vals[c] = 0.0  # No recorded scoring penalty; not a measured assay zero.
                else:
                    vals[c] = float(raw)  # Fail on an undocumented blank/marker.
            score = float(row[score_col])
            assert np.isclose(score, 100 + sum(vals.values())), f"Score reconciliation failed: {cid}"
            signature = json.dumps(vals, sort_keys=True)
            api = api_key(row["API"] if "API" in row and pd.notna(row["API"]) else sheet)
            company = clean(row.Company).upper()
            uid = f"{api}:{company}:{hashlib.sha256(signature.encode()).hexdigest()[:12]}"
            rec = {"score_card_id": cid, "score_unit": uid, "api": api, "score_company": company,
                   "source_sheet": sheet, "source_row": int(idx)+2, "score": score,
                   "quality_loss": 100-score, "application": clean(row["Application Number"]),
                   "component_raw": json.dumps({c:clean(row[c]) for c in components}, sort_keys=True),
                   "dissolution_penalty": -sum(v for c,v in vals.items() if "dissolution" in c.lower()) if any("dissolution" in c.lower() for c in vals) else np.nan,
                   "chemical_penalty": -sum(v for c,v in vals.items() if any(t in c.lower() for t in ["carcinogen", "toxic element"])),
                   "dosage_penalty": -sum(v for c,v in vals.items() if "dosage" in c.lower()),
                   "sterility_penalty": -sum(v for c,v in vals.items() if "sterility" in c.lower()) if any("sterility" in c.lower() for c in vals) else np.nan}
            records.append(rec)
            seen = set()
            for column,value in enumerate(row.iloc[ndc_idx:],ndc_idx+1):
                for match in re.findall(r"\d{4,5}-\d{3,4}-\d{1,2}", clean(value)):
                    n = ndc(match)
                    if n and n not in seen:
                        expanded.append({**rec, "ndc11": n, "ndc9": n[:9],"source_ndc_cell":f"{get_column_letter(column)}{int(idx)+2}"})
                        seen.add(n)
            assert seen, f"No NDC parsed for score card {cid}"
    return pd.DataFrame(records), pd.DataFrame(expanded)


def attach_scores(s: pd.DataFrame, expanded: pd.DataFrame) -> pd.DataFrame:
    groups = {k: g for k,g in expanded.groupby(["api", "ndc9"])}
    rows = []
    wanted = ["score_unit", "score", "quality_loss", "dissolution_penalty", "chemical_penalty", "dosage_penalty", "sterility_penalty"]
    for _, r in s.iterrows():
        g = groups.get((r.api, r.ndc9))
        rec = {"sample_id": r.sample_id, "score_match_status": "not_matched"}
        if g is not None:
            # Product-code matching permits differing package sizes, but never differing score units.
            units = g.score_unit.unique()
            rec["score_candidate_units"] = " ; ".join(sorted(units))
            rec["score_card_ids"] = " ; ".join(sorted(g.score_card_id.unique()))
            rec["score_match_status"] = "matched" if len(units) == 1 else "ambiguous"
            if len(units) == 1:
                first = g.iloc[0]
                rec.update({k:first[k] for k in wanted})
                rec["score_match_status"] = "matched" if np.isclose(r.recorded_drug_score, first.score) else "score_disagrees"
        rows.append(rec)
    return s.merge(pd.DataFrame(rows), on="sample_id", validate="one_to_one")


def facility_links(ndcs: set[str]) -> tuple[pd.DataFrame, dict, dict]:
    links = []
    march = pd.read_excel(FILES["march"], sheet_name="NDC_FEI Mapping")
    for _,r in march.iterrows():
        n, f = ndc(r["NDC"], True), fei_key(r["FEI_NUMBER"])
        if n in ndcs and f:
            links.append({"ndc9":n, "fei":f, "source":"march", "role":"unspecified", "included":True})
    pp = pd.read_csv(FILES["propublica"], dtype=str, low_memory=False)
    for _,r in pp.iterrows():
        n,f = ndc(r.ndc, True), fei_key(r.fei)
        if n in ndcs and f:
            api = clean(r.api_mfr).upper() == "TRUE"
            links.append({"ndc9":n, "fei":f, "source":"propublica", "role":"api_only" if api else "unspecified", "included":not api})
    dm = pd.read_csv(FILES["dailymed"], dtype=str, low_memory=False)
    dm = dm[dm.opr_type.fillna("").str.lower().isin(["manufacture", "fdf manufacture"])]
    for _,r in dm.iterrows():
        n,f = ndc(r.ndc, True), fei_key(r.FEI)
        if n in ndcs and f:
            links.append({"ndc9":n, "fei":f, "source":"dailymed", "role":clean(r.opr_type).lower(), "included":True})
    audit = pd.DataFrame(links).drop_duplicates()
    union = audit[audit.included].groupby("ndc9").fei.apply(lambda x:sorted(set(x))).to_dict()
    old = audit[audit.source.eq("march")].groupby("ndc9").fei.apply(lambda x:sorted(set(x))).to_dict()
    return audit, union, old


def read_observations(apis: list[str]) -> pd.DataFrame:
    d = pd.read_csv(FILES["observations"], dtype={"fei":str})
    d["fei"] = d.fei.map(fei_key)
    d["insp_date"] = pd.to_datetime(d.insp_date)
    d["observation_id"] = d.fei + ":" + d.insp_date.dt.strftime("%Y-%m-%d") + ":" + d.obs_num.astype(str)
    assert not d.observation_id.duplicated().any()
    txt = d.obs_text_clean.fillna("")
    d["sterile_screen"] = txt.str.contains(STERILE)
    d["oral_screen"] = txt.str.contains(ORAL)
    d["screen_dissolution"] = txt.str.contains(DISSOLUTION)
    d["screen_chemical"] = txt.str.contains(CHEMICAL)
    d["patient_risk"] = d.patient_risk_flag_llm.astype(str).str.lower().eq("true")
    d["data_integrity"] = d.data_integrity_flag_llm.astype(str).str.lower().eq("true")
    d["lab"] = d.violation_category.eq("LaboratoryControlsSystem")
    # Combination products are compatible with a mention of a constituent API.
    names = sorted({a for api in apis for a in api.split("; ")})
    pats = {a:re.compile(r"\b" + re.escape(a).replace(r"\ ", r"\s+") + r"\b", re.I) for a in names}
    d["named_apis_screen"] = txt.map(lambda x: "; ".join(a for a,p in pats.items() if p.search(x)))
    return d


def attach_text(s: pd.DataFrame, obs: pd.DataFrame) -> pd.DataFrame:
    # Supplement only the explicitly dated revised metformin source.
    dates = pd.read_excel(FILES["intake"], sheet_name="2024 Testing Data", header=1)
    dates = dates[dates["Sample ID"].notna()].copy()
    dates["intake_date"] = pd.to_datetime(dates["Intake Date"].astype(str).str.replace(r"\.0$", "", regex=True), format="%Y%m%d", errors="coerce")
    assert not dates["Sample ID"].duplicated().any()
    date_map = dates.set_index("Sample ID").intake_date
    s = s.copy()
    s["intake_date"] = s.sample_id.map(date_map)
    s["cutoff_date"] = s.intake_date.fillna(pd.Timestamp("2023-01-01"))
    s["date_basis"] = np.where(s.intake_date.notna(), "observed_metformin_intake", "conservative_2023_boundary")
    grades = pd.read_parquet(FILES["grades"], columns=["fei", "insp_date", "n_oai", "n_vai", "n_nai"])
    grades["fei"] = grades.fei.map(fei_key)
    grades["insp_date"] = pd.to_datetime(grades.insp_date)
    grade = {}
    for key,g in grades.groupby(["fei", "insp_date"]):
        grade[key] = "OAI" if g.n_oai.max() > 0 else ("VAI" if g.n_vai.max() > 0 else ("NAI" if g.n_nai.max() > 0 else "UNKNOWN"))
    groups = {k:g for k,g in obs.groupby("fei")}
    records = []
    for _,r in s.iterrows():
        rec = {"sample_id":r.sample_id, "text_match_status":"no_unique_facility"}
        if not r.fei:
            records.append(rec)
            continue
        g = groups.get(r.fei)
        if g is None or not (g.insp_date < r.cutoff_date).any():
            rec["text_match_status"] = "no_prior_text"
            records.append(rec)
            continue
        snapshot = g.loc[g.insp_date < r.cutoff_date, "insp_date"].max()
        g = g[g.insp_date.eq(snapshot)].copy()
        n_all = len(g)
        keep = pd.Series(True, index=g.index)
        if r.route_screen == "oral_solid":
            keep &= ~(g.sterile_screen & ~g.oral_screen)
        elif r.route_screen == "injectable":
            keep &= ~(g.oral_screen & ~g.sterile_screen)
        relevant = set(r.api.split("; "))
        keep &= g.named_apis_screen.map(lambda x: not x or bool(set(x.split("; ")) & relevant))
        g = g[keep]
        rec.update({"inspection_date":snapshot, "inspection_age_days":int((r.cutoff_date-snapshot).days), "n_obs_all":n_all, "n_obs_compatible":len(g), "fda_class":grade.get((r.fei,snapshot), "UNKNOWN"), "named_product_observations":int(g.named_apis_screen.ne("").sum())})
        rec["text_match_status"] = "matched" if len(g) else "no_compatible_observations"
        for raw,name in [("screen_dissolution","screen_dissolution_share"), ("screen_chemical","screen_chemical_share"), ("patient_risk","patient_risk_share"), ("data_integrity","data_integrity_share"), ("lab","lab_share")]:
            rec[name] = g[raw].mean() if len(g) else np.nan
        records.append(rec)
    return s.merge(pd.DataFrame(records), on="sample_id", validate="one_to_one")


def connected_groups(d: pd.DataFrame) -> pd.Series:
    parent = {}
    def find(x):
        parent.setdefault(x,x)
        if parent[x] != x:
            parent[x] = find(parent[x])
        return parent[x]
    for _,r in d.iterrows():
        a,b = find("F:"+r.fei),find("S:"+r.score_unit)
        if a != b:
            parent[b] = a
    return d.fei.map(lambda x:find("F:"+x))


def cells(s: pd.DataFrame, scored: bool) -> pd.DataFrame:
    x = s[s.text_match_status.eq("matched")].copy()
    if scored:
        x = x[x.score_match_status.eq("matched")].copy()
        x["cluster"] = connected_groups(x)
    else:
        x = x[x.difference_factor.notna()].copy()
        x["cluster"] = x.fei
    if x.empty:
        return x
    keys = ["fei", "product", "api", "form", "strength", "cluster"] + (["score_unit"] if scored else [])
    features = [p for p in [a[1] for a in ASSOCIATIONS] if p in x]
    numeric = list(dict.fromkeys(features + ["inspection_age_days", "n_obs_compatible", "difference_factor"] + (["quality_loss", "dissolution_penalty", "chemical_penalty", "score"] if scored else [])))
    aggregation = {c:"mean" for c in numeric}
    aggregation.update({"n_samples":("sample_id","nunique"), "sample_ids":("sample_id",lambda a:" ; ".join(sorted(set(a)))), "fda_class":("fda_class",lambda a: next(iter(set(a))) if len(set(a)) == 1 else "MIXED"), "has_dissolution_measurement":("has_dissolution_measurement","any"), "has_chemical_measurement":("has_chemical_measurement","any"), "date_basis":("date_basis",lambda a:" ; ".join(sorted(set(a))))})
    spec = {c:(c,fun) for c,fun in aggregation.items() if not isinstance(fun,tuple)}
    spec.update({c:fun for c,fun in aggregation.items() if isinstance(fun,tuple)})
    out = x.groupby(keys,dropna=False).agg(**spec).reset_index()
    out["n_product_facilities"] = out.groupby("product").fei.transform("nunique")
    out["analysis_weight"] = 1 / out.groupby("score_unit").fei.transform("size") if scored else 1.0
    if not scored:
        out["dissolution_percentile"] = (out.groupby("product").difference_factor.rank(method="average") - .5) / out.groupby("product").fei.transform("size")
    return out


def associate(d: pd.DataFrame, target: str, feature: str, label: str, cohort: str) -> dict:
    rec = {"cohort":cohort, "analysis":label, "outcome":target, "feature":feature, "status":"not_estimable"}
    if d.empty or target not in d:
        return rec
    x = d[d.n_product_facilities.ge(2)].copy()
    if target == "dissolution_penalty":
        x = x[x.has_dissolution_measurement]
    elif target == "chemical_penalty":
        x = x[x.has_chemical_measurement]
    x = x.dropna(subset=[target, feature, "inspection_age_days", "n_obs_compatible"])
    # Recheck strata after endpoint-specific missingness/exclusions.
    x = x[x.groupby("product").fei.transform("nunique").ge(2)].copy()
    if "score_unit" in x:
        x["analysis_weight"] = 1 / x.groupby("score_unit").fei.transform("size")
    k = x.cluster.nunique()
    exposed = x.loc[x[feature].gt(0),"cluster"].nunique()
    unexposed = x.loc[x[feature].eq(0),"cluster"].nunique()
    rec.update(n_cells=len(x), n_samples=int(x.n_samples.sum()), n_facilities=x.fei.nunique(), n_clusters=k, n_products=x['product'].nunique(), n_exposed_clusters=exposed, n_zero_exposure_clusters=unexposed)
    if len(x) < 8 or k < 4 or x[target].nunique() < 2 or x[feature].nunique() < 2:
        rec["reason"] = "Too little independent variation or too few clusters"
        return rec
    w = x.analysis_weight.to_numpy()
    f = x[feature].to_numpy(dtype=float)
    sd = np.sqrt(np.average((f-np.average(f,weights=w))**2,weights=w))
    if sd <= 1e-12:
        rec["reason"] = "Constant weighted exposure"
        return rec
    base = pd.DataFrame({"signal":f/sd, "log_age":np.log1p(x.inspection_age_days.to_numpy()), "log_n_obs":np.log1p(x.n_obs_compatible.to_numpy()), "outcome":x[target].to_numpy()}, index=x.index)
    # Weighted within-product transformation = WLS product fixed effects.
    for col in base:
        av = (base[col]*x.analysis_weight).groupby(x['product']).transform("sum") / x.analysis_weight.groupby(x['product']).transform("sum")
        base[col] -= av
    if np.max(np.abs(base.signal)) < 1e-10:
        rec["reason"] = "No exposure variation within product"
        return rec
    X = base[["signal", "log_age", "log_n_obs"]]
    X = X.loc[:,X.abs().max().gt(1e-10)]
    if np.linalg.matrix_rank(X.to_numpy()) < X.shape[1]:
        rec["reason"] = "Collinear within-product design"
        return rec
    # Use a full fixed-effect design so cluster finite-sample corrections include strata degrees of freedom.
    full = pd.concat([x[[feature]].rename(columns={feature:"signal"})/sd,
                      pd.DataFrame({"log_age":np.log1p(x.inspection_age_days), "log_n_obs":np.log1p(x.n_obs_compatible)},index=x.index),
                      pd.get_dummies(x['product'],dtype=float)],axis=1)
    # Remove nuisance covariates constant within strata; the fixed effects absorb them.
    for c in ["log_age", "log_n_obs"]:
        if c not in X:
            full = full.drop(columns=c)
    if len(x) <= np.linalg.matrix_rank(full.to_numpy()) + 1:
        rec["reason"] = "Insufficient residual degrees of freedom"
        return rec
    model = sm.WLS(x[target], full, weights=x.analysis_weight).fit(cov_type="cluster", cov_kwds={"groups":x.cluster,"use_correction":True}, use_t=True)
    beta,se = float(model.params.signal),float(model.bse.signal)
    critical = student_t.ppf(.975,k-1)
    rec.update(status="estimated", effect_per_sd=beta, se=se, ci_low=beta-critical*se, ci_high=beta+critical*se, p_value=float(model.pvalues.signal), exposure_sd=sd, interpretation="Exploratory; small-cluster uncertainty" if k < 20 else "Exploratory association")
    if exposed < 4 and unexposed > 0:
        rec.update(status="descriptive_sparse_exposure", interpretation="Fewer than four exposed clusters; inference withheld",
                   nominal_p_value=rec["p_value"],nominal_ci_low=rec["ci_low"],nominal_ci_high=rec["ci_high"],
                   p_value=np.nan,ci_low=np.nan,ci_high=np.nan)
    return rec


def prediction_comparison(d: pd.DataFrame, cohort: str, bootstraps: int) -> tuple[dict,pd.DataFrame]:
    rec = {"cohort":cohort,"status":"not_estimable"}
    if d.empty:
        return rec,pd.DataFrame()
    x = d[d.n_product_facilities.ge(2)].copy().reset_index(drop=True)
    x["analysis_weight"] = 1 / x.groupby("score_unit").fei.transform("size")
    y = x.quality_loss.gt(0).astype(int).to_numpy()
    groups = x.cluster.to_numpy()
    rec.update(n_cells=len(x),n_facilities=x.fei.nunique(),n_clusters=x.cluster.nunique())
    if len(x) < 30 or x.cluster.nunique() < 8 or len(np.unique(y)) < 2:
        rec["reason"] = "Too few independent cells/clusters or only one outcome class"
        return rec,pd.DataFrame()
    x["log_age"] = np.log1p(x.inspection_age_days)
    x["log_n_obs"] = np.log1p(x.n_obs_compatible)
    base_num = ["log_age", "log_n_obs"]
    prob = {"baseline":np.full(len(x),np.nan),"baseline_plus_text":np.full(len(x),np.nan)}
    folds = np.zeros(len(x),dtype=int)
    for fold,(train,test) in enumerate(GroupKFold(n_splits=5).split(x,y,groups),1):
        folds[test] = fold
        for name,nums in [("baseline",base_num),("baseline_plus_text",base_num+LLM_FIELDS)]:
            if len(np.unique(y[train])) < 2:
                prob[name][test] = np.average(y[train],weights=x.analysis_weight.iloc[train])
                continue
            transform = ColumnTransformer([("cat",OneHotEncoder(handle_unknown="ignore"),["product","fda_class"]), ("num",StandardScaler(),nums)])
            model = make_pipeline(transform,LogisticRegression(C=1.0,max_iter=2000,random_state=42))
            model.fit(x.iloc[train],y[train],logisticregression__sample_weight=x.analysis_weight.iloc[train])
            prob[name][test] = model.predict_proba(x.iloc[test])[:,1]
    w = x.analysis_weight.to_numpy()
    metrics = {}
    for name,p in prob.items():
        metrics[name+"_auc"] = roc_auc_score(y,p,sample_weight=w)
        metrics[name+"_brier"] = brier_score_loss(y,p,sample_weight=w)
    delta = metrics["baseline_plus_text_auc"]-metrics["baseline_auc"]
    unique = np.unique(groups)
    positions = {g:np.flatnonzero(groups==g) for g in unique}
    rng = np.random.default_rng(20261006)
    diffs = []
    for _ in range(bootstraps):
        idx = np.concatenate([positions[g] for g in rng.choice(unique,size=len(unique),replace=True)])
        if len(np.unique(y[idx])) > 1:
            diffs.append(roc_auc_score(y[idx],prob["baseline_plus_text"][idx],sample_weight=w[idx])-roc_auc_score(y[idx],prob["baseline"][idx],sample_weight=w[idx]))
    interval = np.quantile(diffs,[.025,.975]) if diffs else [np.nan,np.nan]
    rec.update(status="estimated",**metrics,auc_difference=delta,delta_ci_low=float(interval[0]),delta_ci_high=float(interval[1]),valid_bootstraps=len(diffs),uncertainty="Paired connected-group bootstrap of fixed OOF predictions; conditional on trained models")
    pred = x[["fei","product","score_unit","cluster","quality_loss","analysis_weight","fda_class"]].copy()
    pred["cohort"],pred["fold"],pred["recorded_penalty"] = cohort,folds,y
    for name,p in prob.items():
        pred[name+"_probability"] = p
    assert pred.groupby("cluster").fold.nunique().max() == 1
    return rec,pred


def markdown_table(d: pd.DataFrame, cols: list[str]) -> str:
    d = d.reindex(columns=cols)
    rows = ["| " + " | ".join(cols) + " |", "| " + " | ".join(["---"]*len(cols)) + " |"]
    for _,r in d.iterrows():
        vals = []
        for v in r:
            if pd.isna(v):
                vals.append("")
            elif isinstance(v,(float,np.floating)):
                vals.append(f"{v:.3f}")
            else:
                vals.append(str(v).replace("|","/"))
        rows.append("| " + " | ".join(vals) + " |")
    return "\n".join(rows)


def save_figure(out: Path, results: pd.DataFrame):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    d = results[results.cohort.eq("all_source_unique")].set_index("analysis")
    panels = [
        ("Overall recorded score loss", [ASSOCIATIONS[i][2] for i in [2,3,4]], ["Patient-risk flag", "Data-integrity flag", "Laboratory-system share"], 1, "Score-loss points per 1 SD of text exposure"),
        ("Recorded component penalties", [ASSOCIATIONS[i][2] for i in [0,1]], ["Dissolution mention", "Chemical/impurity mention"], 1, "Penalty points per 1 SD of text exposure"),
        ("Raw dissolution within exact product", [ASSOCIATIONS[5][2]], ["Dissolution mention"], 100, "Percentile points per 1 SD of text exposure"),
    ]
    fig,axes = plt.subplots(3,1,figsize=(10.4,7.6),gridspec_kw={"height_ratios":[3,2,1.5]})
    for ax,(title,names,labels,scale,xlabel) in zip(axes,panels):
        for i,name in enumerate(names):
            row = d.loc[name]
            if pd.notna(row.get("ci_low")):
                ax.plot([scale*row.ci_low,scale*row.ci_high],[i,i],color="#526779",lw=2)
                ax.scatter([scale*row.effect_per_sd],[i],color="#176490",s=42,zorder=3)
        ax.axvline(0,color="#888888",lw=1,ls="--")
        ax.set_yticks(range(len(labels)),labels)
        ax.invert_yaxis()
        ax.set_ylim(len(labels)-.5,-.5)
        ax.set_title(title,loc="left",fontsize=11,fontweight="bold")
        ax.set_xlabel(xlabel,fontsize=10)
        ax.spines[["top","right","left"]].set_visible(False)
        ax.tick_params(axis="y",length=0)
    fig.suptitle("Inspection text and Valisure quality measurements",fontsize=14,x=.04,ha="left")
    fig.text(.04,.025,"Main cohort; product-adjusted exploratory effects with 95% cluster intervals.\nKeyword screens and historical product-to-facility attribution require adjudication.",fontsize=9,color="#555555")
    fig.tight_layout(rect=(0,.075,1,.94),h_pad=1.4)
    fig.savefig(out/"association_effects.png",dpi=180)
    plt.close(fig)


def write_report(out: Path, samples: pd.DataFrame, scores: pd.DataFrame, assays: pd.DataFrame, flow: pd.DataFrame, coverage: pd.DataFrame, results: pd.DataFrame, predictions: pd.DataFrame):
    status = samples.score_match_status.value_counts().to_dict()
    primary = results[results.cohort.eq("all_source_unique")]
    text = f"""# Valisure and inspection text: first exploratory results

This run uses all 13 API worksheets in both DoD workbooks, including the separate ampicillin–sulbactam labels. It evaluates recorded laboratory quality and scoring penalties, not patient outcomes. The analysis choices are in the adjacent README. Neither keyword screens nor facility attribution are yet expert-validated for the tested products.

## Interpretation

The main cohort does not establish a reproducible text/quality association: none of the six declared tests survives Holm correction, and adding the three existing LLM features does not improve the paired held-out prediction. The patient-risk flag has a positive nominal association with worse recorded quality and is a candidate for independent, mechanism-specific follow-up. The dissolution and chemical mention screens do not show a clear association in the main cohort. These results do not establish that the underlying defect mechanisms are absent or clinically irrelevant.

![Exploratory associations and cluster intervals](association_effects.png)

## Source reconciliation

- Testing workbook: {len(samples)} unique sample IDs, {samples.source_sheet.nunique()} worksheets, {samples.api.nunique()} distinct API labels.
- Scoring workbook: {len(scores)} source score rows and {scores.score_unit.nunique()} conservatively shared score units. All source totals reconcile to 100 plus component penalties.
- Source scores range from {scores.score.min():g} to {scores.score.max():g}; they are not clipped to 0–100.
- Scoring match statuses: `{status}`. Score disagreements/ambiguous matches are excluded from score models, but retained in audit files.
- NDC parsing: `{samples.ndc_parse_status.value_counts().to_dict()}`. One malformed package segment can retain an unambiguous product code; the raw identifier is preserved and the package code is not silently repaired.
- Assay status counts: `{assays.status.value_counts().to_dict()}`. Censored and absent values are distinct; neither is substituted with zero.
- Timing: {samples.intake_date.notna().sum()} metformin samples have observed intake dates. Other samples use the conservative January 1, 2023 boundary. This cannot establish lot manufacture time or future prediction.

## Cohort flow

{markdown_table(flow,list(flow.columns))}

Main linkage requires one candidate FEI in the full cross-source union before selecting text-covered plants. DailyMed manufacture/FDF roles are included; ProPublica explicitly API-only links are excluded. Other links can still have uncertain roles and dates. The March-only analysis deliberately admits cross-source disagreements and is a sensitivity analysis.

## Coverage by API

{markdown_table(coverage,list(coverage.columns))}

## All six declared associations in the main cohort

Effects are outcome units per weighted between-cell SD of the exposure, estimated within exact product and adjusted for inspection age/observation count. Scores have one total weight per shared score unit. Connected FEI/score-unit clusters account for both shared plants and repeated scores. Raw dissolution uses within-product percentiles and FEI clustering. The numeric raw value is never given an unverified clinical/regulatory cutoff.

{markdown_table(primary,["analysis","status","n_cells","n_facilities","n_clusters","n_exposed_clusters","effect_per_sd","ci_low","ci_high","p_value","p_holm"])}

All sensitivities and skipped models are in `association_results.csv`. A small p-value does not validate attribution or the keyword mechanism. Lack of an association does not establish clinical equivalence. Multiple products and cohorts have already been examined, so these are exploratory results even with multiplicity correction.

The dated metformin raw-dissolution analysis has only one facility with a positive dissolution mention screen. Its nominal association is not replicated across exposed facilities. Such sparse-exposure models retain descriptive coefficients and nominal outputs in the CSV, but their inferential p-values/intervals are withheld from the interpreted results. The main raw-dissolution analysis has only five exposed facilities, which is also a substantial precision and generalizability limit.

Reading the candidate text confirms why semantic adjudication matters. Observation `3005406526:2019-06-07:6` discusses dissolution-bath qualification/calibration, not a confirmed product dissolution failure. Observation `3004554612:2022-12-09:3` discusses product dissolution OOS results and market batches, but the product identity is redacted. Observation `3002809586:2022-05-09:9` mixes several products and mechanisms and includes a challenged dissolution-result invalidation. A mention count cannot distinguish these situations or establish that the tested API was affected. These examples are source-review notes, not expert-validated new labels, and no additional association was selected after reading them.

## Paired held-out prediction comparison

The endpoint is any recorded scoring penalty, not patient harm. Both models use identical observations and folds. The baseline contains exact product, FDA class (unknown retained), inspection age and observation count; the extension adds the three existing LLM patient-risk, data-integrity and laboratory-system shares. Every connected score/site group stays in one fold.

{markdown_table(predictions,["cohort","status","n_cells","n_facilities","n_clusters","baseline_auc","baseline_plus_text_auc","auc_difference","delta_ci_low","delta_ci_high","baseline_brier","baseline_plus_text_brier"])}

Bootstrap intervals resample groups of fixed out-of-fold predictions. They quantify conditional test-sample uncertainty and do not include all model-training or specification-selection uncertainty. Results need an untouched validation cohort before a predictive claim.

## What is ready for review

`samples_all_source_unique.csv` retains every sample, score match, candidate sites, dates and inclusion reasons. `assays_long.csv` preserves raw values, units in assay names, censoring, and source cells. `scorecards.csv` and `scorecard_ndcs.csv` reconcile scoring provenance. `facility_link_audit.csv` lists the contributing source/role for each product/site link. `annotation_blind.csv` provides the full observation text, candidate mechanism/route screens, and empty expert-review fields without laboratory outcomes. `annotation_priority_blind.csv` restricts this to inspections actually linked in the main cohort and places candidate dissolution/chemical mentions first; its ordering uses no assay outcomes.

Next steps are expert adjudication of those mechanisms, historical finished-dose/product linkage and assay documentation. The broad facility-text association is not the same as demonstrating a matching defect in the tested product. The original manuscripts, workbooks and patient-outcome scripts were not modified.
"""
    (out/"RESULTS.md").write_text(text)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output-dir",type=Path,default=HERE/"outputs/20261006")
    parser.add_argument("--bootstraps",type=int,default=2000)
    args = parser.parse_args()
    out = args.output_dir
    out.mkdir(parents=True,exist_ok=True)
    samples,assays = read_samples()
    scores,expanded = read_scores()
    samples = attach_scores(samples,expanded)
    audit,union,march = facility_links(set(samples.ndc9))
    obs = read_observations(sorted(samples.api.unique()))
    sources = {name:{"path":str(path.relative_to(ROOT)),"sha256":hashlib.sha256(path.read_bytes()).hexdigest()} for name,path in FILES.items()}
    sources["specification"] = {"sha256":hashlib.sha256((HERE/"README.md").read_bytes()).hexdigest()}
    sources["run"] = {"script_sha256":hashlib.sha256(Path(__file__).read_bytes()).hexdigest(),"bootstraps":args.bootstraps,"bootstrap_seed":20261006,"model_seed":42,"pandas_version":pd.__version__,"numpy_version":np.__version__}
    (out/"source_manifest.json").write_text(json.dumps(sources,indent=2))
    assays.to_csv(out/"assays_long.csv",index=False)
    scores.to_csv(out/"scorecards.csv",index=False)
    expanded.to_csv(out/"scorecard_ndcs.csv",index=False)
    audit.to_csv(out/"facility_link_audit.csv",index=False)
    blind_cols = ["observation_id","fei","insp_date","obs_num","obs_text_clean","violation_category","severity_tier","patient_risk_flag_llm","data_integrity_flag_llm","screen_dissolution","screen_chemical","sterile_screen","oral_screen","named_apis_screen"]
    blind = obs[blind_cols].copy()
    for c in ["expert_mechanism","expert_named_product","expert_route_form_strength","expert_confirmed_defect","expert_distributed_product","expert_historical_period","expert_supporting_span","reviewer","adjudication_notes"]:
        blind[c] = ""
    blind.to_csv(out/"annotation_blind.csv",index=False)
    results,flows,pred_results,pred_rows = [],[],[],[]
    main_samples = None
    for cohort,mapping in [("all_source_unique",union),("march_only_unique",march),("dated_metformin",union)]:
        s = samples.copy()
        s["candidate_feis"] = s.ndc9.map(lambda n:" ; ".join(union.get(n,[])))
        s["n_candidate_feis"] = s.ndc9.map(lambda n:len(union.get(n,[])))
        s["march_candidate_feis"] = s.ndc9.map(lambda n:" ; ".join(march.get(n,[])))
        s["fei"] = s.ndc9.map(lambda n:mapping.get(n,[""])[0] if len(mapping.get(n,[])) == 1 else "")
        s = attach_text(s,obs)
        if cohort == "dated_metformin":
            s = s[s.intake_date.notna()].copy()
        if cohort == "all_source_unique":
            main_samples = s
        s.to_csv(out/f"samples_{cohort}.csv",index=False)
        score_cells,raw_cells = cells(s,True),cells(s,False)
        score_cells.to_csv(out/f"score_cells_{cohort}.csv",index=False)
        raw_cells.to_csv(out/f"dissolution_cells_{cohort}.csv",index=False)
        matched = s[s.text_match_status.eq("matched")]
        flows.append({"cohort":cohort,"source_samples":len(s),"unique_site_samples":int(s.fei.ne("").sum()),"prior_compatible_text_samples":len(matched),"text_facilities":matched.fei.nunique(),"scored_cells":len(score_cells),"dissolution_cells":len(raw_cells)})
        for target,feature,label in ASSOCIATIONS:
            results.append(associate(raw_cells if target == "dissolution_percentile" else score_cells,target,feature,label,cohort))
        metrics,pred = prediction_comparison(score_cells,cohort,args.bootstraps)
        pred_results.append(metrics)
        if not pred.empty:
            pred_rows.append(pred)
    res = pd.DataFrame(results)
    res["p_holm"] = np.nan
    for _,g in res.groupby("cohort"):
        valid = g[g.get("p_value",pd.Series(index=g.index,dtype=float)).notna()]
        if len(valid):
            # Keep the declared six-test family even when some endpoints are not estimable.
            p = np.r_[valid.p_value.to_numpy(),np.ones(len(ASSOCIATIONS)-len(valid))]
            res.loc[valid.index,"p_holm"] = multipletests(p,method="holm")[1][:len(valid)]
    flow = pd.DataFrame(flows)
    pred_summary = pd.DataFrame(pred_results)
    res.to_csv(out/"association_results.csv",index=False)
    flow.to_csv(out/"cohort_flow.csv",index=False)
    pred_summary.to_csv(out/"prediction_comparison.csv",index=False)
    if pred_rows:
        pd.concat(pred_rows,ignore_index=True).to_csv(out/"oof_predictions.csv",index=False)
    coverage = []
    for api,g in main_samples.groupby("api"):
        m = g[g.text_match_status.eq("matched")]
        coverage.append({"api":api,"samples":len(g),"unique_site_samples":int(g.fei.ne("").sum()),"prior_text_samples":len(m),"prior_text_facilities":m.fei.nunique(),"raw_dissolution_samples":int(g.has_dissolution_measurement.sum()),"chemical_measured_samples":int(g.has_chemical_measurement.sum())})
    coverage = pd.DataFrame(coverage)
    coverage.to_csv(out/"api_coverage.csv",index=False)
    prior = main_samples[main_samples.text_match_status.eq("matched")][["fei","inspection_date"]].drop_duplicates().rename(columns={"inspection_date":"insp_date"})
    priority = blind.merge(prior,on=["fei","insp_date"],how="inner",validate="many_to_one")
    priority.sort_values(["screen_dissolution","screen_chemical","fei","insp_date","obs_num"],ascending=[False,False,True,True,True]).to_csv(out/"annotation_priority_blind.csv",index=False)
    save_figure(out,res)
    write_report(out,main_samples,scores,assays,flow,coverage,res,pred_summary)
    print(flow.to_string(index=False))
    print(res[res.cohort.eq("all_source_unique")].to_string(index=False))
    print(pred_summary.to_string(index=False))
    print(f"Saved {out/'RESULTS.md'}")


if __name__ == "__main__":
    main()
